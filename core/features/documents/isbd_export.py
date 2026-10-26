"""
iSBD export — fills a source iSBD PPTX with editor data and renders a
native Gantt timeline. Phase codes/colors come from the shared domain
contract (shared/domain_constants.json) so the exported deck always
matches the frontend palette.
"""
import tempfile

from core.shared.domain import PHASE_COLORS_RGB, PHASE_LABELS, PHASE_NAME_TO_CODE
from core.shared.field_index import isbd_key_for
from core.features.documents.isbd_maps import (
    RISK_COLS,
    SLIDE_TITLE_MAP,
    SOLUTION_ROW_KEYS,
)
from core.features.documents.isbd_extract import get_slide_title, table_header_text

_MONTH_ABBR = ['Jan', 'Feb', 'Mar', 'Apr', 'May', 'Jun', 'Jul', 'Aug', 'Sep', 'Oct', 'Nov', 'Dec']
_DEFAULT_BAR_COLOR = '6F6F6F'


def _phase_color(code):
    return PHASE_COLORS_RGB.get(code, _DEFAULT_BAR_COLOR)


def _set_cell_text(cell, text):
    tf = cell.text_frame
    while len(tf.paragraphs) > 1:
        p = tf.paragraphs[-1]._p
        p.getparent().remove(p)
    para = tf.paragraphs[0]
    if para.runs:
        para.runs[0].text = str(text)
        for extra in list(para.runs)[1:]:
            extra._r.getparent().remove(extra._r)
    else:
        para.add_run().text = str(text)


def _rebuild_table_rows(table, col_keys, row_data):
    from copy import deepcopy
    if not table.rows or not row_data:
        return
    template_row = table.rows[-1]
    existing_data_rows = list(table.rows)[1:]
    for row in existing_data_rows:
        tr = row._tr
        tr.getparent().remove(tr)
    tbl = table._tbl
    for record in row_data:
        new_tr = deepcopy(template_row._tr)
        cells = new_tr.findall('.//{http://schemas.openxmlformats.org/drawingml/2006/main}tc')
        for i, key in enumerate(col_keys):
            if i < len(cells):
                cell_tf = cells[i].find(
                    '{http://schemas.openxmlformats.org/drawingml/2006/main}txBody'
                )
                if cell_tf is not None:
                    for r in cell_tf.findall(
                        './/{http://schemas.openxmlformats.org/drawingml/2006/main}r'
                    ):
                        t_elem = r.find(
                            '{http://schemas.openxmlformats.org/drawingml/2006/main}t'
                        )
                        if t_elem is not None:
                            t_elem.text = str(record.get(key, ''))
                            break
        tbl.append(new_tr)


def _fill_slide_kv(table, isbd_data, slide_type):
    for row in table.rows:
        cells = list(row.cells)
        # Single-column table (cover format): "Label: Value" in one cell
        if len(cells) == 1 and ':' in cells[0].text:
            cell_text = cells[0].text.strip()
            label, _, old_val = cell_text.partition(':')
            label = label.strip()
            field_key = isbd_key_for(label, slide_type)
            if field_key and field_key in isbd_data:
                new_val = isbd_data[field_key]
                _set_cell_text(cells[0], f'{label}: {new_val}')
        # Two-column table: col 0 = label, col 1 = value
        if len(cells) >= 2:
            label = cells[0].text.strip()
            field_key = isbd_key_for(label, slide_type)
            if field_key and field_key in isbd_data:
                _set_cell_text(cells[1], isbd_data[field_key])
        # Four-column table: col 2 = label, col 3 = value
        if len(cells) >= 4:
            label = cells[2].text.strip()
            field_key = isbd_key_for(label, slide_type)
            if field_key and field_key in isbd_data:
                _set_cell_text(cells[3], isbd_data[field_key])


def _render_gantt_on_slide(slide, timelines, Emu, Pt, RGBColor, PP_ALIGN, MSO_ANCHOR):
    """Replace the existing image on the proposed_timeline slide with native Gantt bars."""
    from pptx.enum.shapes import MSO_SHAPE

    # Remove existing picture shapes (the static WMF/PNG chart)
    sp_tree = slide.shapes._spTree
    to_remove = []
    for shape in slide.shapes:
        if shape.shape_type is not None and shape.shape_type == 13:  # MSO_SHAPE_TYPE.PICTURE
            to_remove.append(shape._element)
    for el in to_remove:
        sp_tree.remove(el)

    # Convert timelines to lane data (mirrors frontend timelinesToLanes)
    lanes = []
    for tl in timelines:
        phase_map = tl.get('timeline', {})
        segments = []
        offset = 0
        for phase_name, phase_data in phase_map.items():
            dur = phase_data if isinstance(phase_data, dict) else {'duration': phase_data, 'unit': 'months'}
            months = dur.get('duration', 2)
            if dur.get('unit') == 'weeks':
                months = max(1, round(months / 4.33))
            code = PHASE_NAME_TO_CODE.get(phase_name, phase_name[:2].upper())
            if months > 0:
                segments.append({'code': code, 'start': offset, 'span': months, 'phase': phase_name})
                offset += months
        lanes.append({
            'name': tl.get('name', 'Timeline'),
            'start_date': tl.get('project_start_date', ''),
            'total_months': offset or tl.get('total_duration', 0),
            'segments': segments,
        })

    if not lanes:
        return

    # ── Layout constants (EMU) — chart region matching ref slide 5 Picture 4 ──
    chart_left = 2249263
    chart_top = 1847261
    chart_width = 9532980
    chart_height = 2806933

    # Label column to the left of the chart bars
    label_width = 1700000
    bar_left = chart_left + label_width
    bar_width = chart_width - label_width

    # Calculate grid
    max_months = max(l['total_months'] for l in lanes) or 1
    col_w = bar_width // max_months
    lane_count = len(lanes)
    axis_height = 280000  # space for month labels
    lane_height = (chart_height - axis_height) // max(lane_count, 1)
    bar_h = max(180000, lane_height - 80000)

    # ── Draw month axis labels ──
    for m in range(max_months):
        # Try to use real month names from the first lane's start date
        lbl = f'M{m + 1}'
        sd = lanes[0].get('start_date', '')
        if sd:
            try:
                from datetime import datetime
                dt = datetime.fromisoformat(sd.replace('Z', '+00:00'))
                # Add m months
                month = (dt.month - 1 + m) % 12
                year = dt.year + (dt.month - 1 + m) // 12
                lbl = f"{_MONTH_ABBR[month]} {str(year)[-2:]}"
            except Exception:
                pass

        shape = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Emu(bar_left + m * col_w), Emu(chart_top),
            Emu(col_w), Emu(axis_height)
        )
        shape.fill.background()
        shape.line.fill.background()
        tf = shape.text_frame
        tf.word_wrap = False
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        run = p.add_run()
        run.text = lbl
        run.font.size = Pt(7)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0x57, 0x60, 0x6A)
        tf.auto_size = None
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0

    # ── Draw a thin separator under axis ──
    line_shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Emu(bar_left), Emu(chart_top + axis_height),
        Emu(bar_width), Emu(12000)
    )
    line_shape.fill.solid()
    line_shape.fill.fore_color.rgb = RGBColor(0xD0, 0xD7, 0xDE)
    line_shape.line.fill.background()

    # ── Draw lanes ──
    for lane_idx, lane in enumerate(lanes):
        lane_y = chart_top + axis_height + 20000 + lane_idx * lane_height

        # Lane label
        label_shape = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Emu(chart_left), Emu(lane_y),
            Emu(label_width - 40000), Emu(lane_height - 20000)
        )
        label_shape.fill.background()
        label_shape.line.fill.background()
        tf = label_shape.text_frame
        tf.word_wrap = True
        tf.auto_size = None
        tf.margin_left = Emu(36000)
        tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.LEFT
        run = p.add_run()
        run.text = lane['name']
        run.font.size = Pt(8)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0x24, 0x29, 0x2F)
        # Duration subtitle
        sub = tf.add_paragraph()
        sub.alignment = PP_ALIGN.LEFT
        sub_run = sub.add_run()
        sub_run.text = f"{lane['total_months']}mo"
        sub_run.font.size = Pt(6)
        sub_run.font.color.rgb = RGBColor(0x65, 0x6D, 0x76)

        # Segments (Gantt bars)
        seg_y = lane_y + (lane_height - bar_h) // 2
        for seg in lane['segments']:
            x = bar_left + seg['start'] * col_w
            w = seg['span'] * col_w
            color_hex = _phase_color(seg['code'])
            r_c, g_c, b_c = int(color_hex[:2], 16), int(color_hex[2:4], 16), int(color_hex[4:], 16)

            bar_shape = slide.shapes.add_shape(
                MSO_SHAPE.ROUNDED_RECTANGLE,
                Emu(x + 8000), Emu(seg_y),
                Emu(w - 16000), Emu(bar_h)
            )
            bar_shape.fill.solid()
            bar_shape.fill.fore_color.rgb = RGBColor(r_c, g_c, b_c)
            bar_shape.line.fill.background()
            # Minimal corner rounding
            bar_shape.adjustments[0] = 0.05

            # Label inside bar
            tf = bar_shape.text_frame
            tf.word_wrap = False
            tf.auto_size = None
            tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.CENTER
            run = p.add_run()
            # Show phase code + duration if wide enough
            if seg['span'] >= 2:
                run.text = f"{seg['code']} · {seg['span']}mo"
            else:
                run.text = seg['code']
            run.font.size = Pt(7)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    # ── Legend row below chart ──
    used_codes = set()
    for lane in lanes:
        for seg in lane['segments']:
            used_codes.add(seg['code'])
    legend_items = [(c, _phase_color(c)) for c in PHASE_COLORS_RGB if c in used_codes]
    if legend_items:
        legend_y = chart_top + chart_height + 80000
        legend_x = chart_left
        swatch_w = 140000
        item_w = 900000
        for i, (code, hex_color) in enumerate(legend_items):
            x = legend_x + i * item_w
            r_c, g_c, b_c = int(hex_color[:2], 16), int(hex_color[2:4], 16), int(hex_color[4:], 16)
            # Swatch
            sw = slide.shapes.add_shape(
                MSO_SHAPE.RECTANGLE,
                Emu(x), Emu(legend_y),
                Emu(swatch_w), Emu(swatch_w)
            )
            sw.fill.solid()
            sw.fill.fore_color.rgb = RGBColor(r_c, g_c, b_c)
            sw.line.fill.background()
            # Label
            lbl_shape = slide.shapes.add_shape(
                MSO_SHAPE.RECTANGLE,
                Emu(x + swatch_w + 30000), Emu(legend_y - 20000),
                Emu(item_w - swatch_w - 60000), Emu(swatch_w + 40000)
            )
            lbl_shape.fill.background()
            lbl_shape.line.fill.background()
            tf = lbl_shape.text_frame
            tf.word_wrap = False
            tf.auto_size = None
            tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
            p = tf.paragraphs[0]
            run = p.add_run()
            run.text = PHASE_LABELS.get(code, code)
            run.font.size = Pt(6)
            run.font.color.rgb = RGBColor(0x57, 0x60, 0x6A)


def export_isbd_pptx(source_path, isbd_data, timelines=None):
    """Clone the source iSBD PPTX, fill tables/kv fields from isbd_data,
    render the Gantt timeline, and return the temp output path."""
    from pptx import Presentation
    from pptx.util import Emu, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
    prs = Presentation(source_path)

    for slide in prs.slides:
        title_text = get_slide_title(slide)
        slide_type = SLIDE_TITLE_MAP.get(title_text.lower().strip()) if title_text else None

        for shape in slide.shapes:
            if not shape.has_table:
                continue
            table = shape.table
            header = table_header_text(table)

            _fill_slide_kv(table, isbd_data, slide_type)

            if slide_type == 'compliance_review':
                if 'compliance item' in header and 'Compliance' in isbd_data:
                    _rebuild_table_rows(table, ['compliance_item', 'status', 'comments'], isbd_data['Compliance'])
                elif 'review item' in header and 'Solution Reviews' in isbd_data:
                    _rebuild_table_rows(table, ['review_item', 'status', 'comments'], isbd_data['Solution Reviews'])

            elif slide_type == 'risk_assessment':
                if ('rca question' in header or '#' in header) and 'rca_rows' in isbd_data:
                    _rebuild_table_rows(table, ['number', 'question', 'assessment', 'reason'], isbd_data['rca_rows'])
                elif 'description' in header and 'probability' in header and 'risk_rows' in isbd_data:
                    _rebuild_table_rows(table, RISK_COLS, isbd_data['risk_rows'])

            elif slide_type == 'staffing_approvals':
                if 'location' in header and 'band mix' in header and 'Delivery Locations' in isbd_data:
                    _rebuild_table_rows(table, ['location', 'band_mix', 'staffing_pct', 'approver', 'approval_date'], isbd_data['Delivery Locations'])
                elif 'description' in header and 'probability' in header and 'Key Resourcing Risks' in isbd_data:
                    _rebuild_table_rows(table, RISK_COLS, isbd_data['Key Resourcing Risks'])

            elif slide_type == 'deal_profile_summary':
                if 'service line' in header and 'IBM Consulting Lines' in isbd_data:
                    _rebuild_table_rows(table, ['col_0', 'service_line', 'sub_line', 'tcv_usd'], isbd_data['IBM Consulting Lines'])

        if slide_type == 'solution_history' and 'submissions' in isbd_data:
            for shape in slide.shapes:
                if shape.has_table and len(shape.table.columns) >= 3:
                    table = shape.table
                    submissions = isbd_data['submissions']
                    deltas = isbd_data.get('deltas', {})
                    rows_raw = list(table.rows)
                    for row in rows_raw:
                        cells = list(row.cells)
                        if not cells or not cells[0].text.strip():
                            continue
                        row_key = SOLUTION_ROW_KEYS.get(cells[0].text.strip().lower())
                        if not row_key:
                            continue
                        for i, sub in enumerate(submissions):
                            if i + 1 < len(cells):
                                _set_cell_text(cells[i + 1], sub.get('values', {}).get(row_key, ''))
                        if len(cells) > len(submissions) + 1:
                            _set_cell_text(cells[-1], deltas.get(row_key, ''))
                    break

        # ── Proposed timeline: replace static image with live Gantt bars ──
        if slide_type == 'proposed_timeline' and timelines:
            _render_gantt_on_slide(slide, timelines, Emu, Pt, RGBColor, PP_ALIGN, MSO_ANCHOR)

    tmp = tempfile.NamedTemporaryFile(suffix='.pptx', delete=False)
    prs.save(tmp.name)
    tmp.close()
    return tmp.name
