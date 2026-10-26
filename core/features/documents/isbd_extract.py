"""
iSBD extraction — parses tables & text frames from an iSBD PPTX into
structured key-value data the editor can populate.

Shared table helpers (get_slide_title, table_header_text, table_to_rows)
also serve the export pipeline in isbd_export.py.
"""
from core.features.documents.isbd_maps import (
    ISBD_MIN_TITLE_MATCHES,
    ISBD_SIGNATURE_TITLES,
    RISK_COLS,
    SLIDE_TITLE_MAP,
    SOLUTION_ROW_KEYS,
)
from core.shared.field_index import isbd_key_for


def get_slide_title(slide):
    for shape in slide.shapes:
        name_lower = (shape.name or '').lower()
        if shape.has_text_frame and 'title' in name_lower:
            text = shape.text_frame.text.strip()
            if text:
                return text
    for shape in slide.shapes:
        if shape.has_text_frame:
            text = shape.text_frame.text.strip()
            if text and len(text) < 100:
                return text
    return None


def table_header_text(table):
    if not table.rows:
        return ''
    return '|'.join(c.text.strip().lower() for c in table.rows[0].cells)


def table_to_rows(table, col_keys):
    rows = []
    for row in list(table.rows)[1:]:
        cells = [c.text.strip() for c in row.cells]
        if not any(cells):
            continue
        rows.append({key: cells[i] if i < len(cells) else '' for i, key in enumerate(col_keys)})
    return rows


def _extract_kv_from_slide(slide, fields, slide_type):
    def put(label, value):
        key = isbd_key_for(label, slide_type)
        if key and value and value != label:
            fields[key] = value

    for shape in slide.shapes:
        if shape.has_text_frame and not shape.has_table:
            text = shape.text_frame.text.strip()
            if ':' in text and '\n' not in text:
                label, _, value = text.partition(':')
                put(label.strip(), value.strip())
        if shape.has_table:
            for row in shape.table.rows:
                cells = [c.text.strip() for c in row.cells]
                if not cells:
                    continue
                if len(cells) >= 2:
                    put(cells[0], cells[1])
                if len(cells) >= 4:
                    put(cells[2], cells[3])
                if len(cells) == 1 and ':' in cells[0]:
                    label, _, val = cells[0].partition(':')
                    put(label.strip(), val.strip())


def _extract_solution_history(slide, fields):
    for shape in slide.shapes:
        if not shape.has_table:
            continue
        table = shape.table
        if len(table.columns) < 3:
            continue
        rows_raw = [[c.text.strip() for c in row.cells] for row in table.rows]
        if not rows_raw:
            continue
        header_idx = 0
        for i, row in enumerate(rows_raw):
            joined = ' '.join(row).lower()
            if 'difference' in joined or any('/' in cell for cell in row[1:] if cell):
                header_idx = i
                break
        header = rows_raw[header_idx]
        num_subs = max(len(header) - 2, 1)
        submissions = [
            {'date': header[i + 1] if i + 1 < len(header) - 1 else '', 'values': {}}
            for i in range(num_subs)
        ]
        deltas = {}
        for row in rows_raw[header_idx + 1:]:
            if not row or not row[0]:
                continue
            row_key = SOLUTION_ROW_KEYS.get(row[0].lower().strip())
            if not row_key:
                continue
            for i, sub in enumerate(submissions):
                if i + 1 < len(row):
                    sub['values'][row_key] = row[i + 1]
            if len(row) > num_subs + 1:
                deltas[row_key] = row[-1]
        if any(s['values'] for s in submissions):
            fields['submissions'] = submissions
            fields['deltas'] = deltas
        return


def _extract_compliance_review(slide, fields):
    tables = [s.table for s in slide.shapes if s.has_table]
    compliance_cols = ['compliance_item', 'status', 'comments']
    review_cols = ['review_item', 'status', 'comments']
    for table in tables:
        header = table_header_text(table)
        if 'compliance item' in header:
            rows = table_to_rows(table, compliance_cols)
            if rows:
                fields['Compliance'] = rows
        elif 'review item' in header or 'solution review' in header:
            rows = table_to_rows(table, review_cols)
            if rows:
                fields['Solution Reviews'] = rows


def _extract_risk_assessment(slide, fields):
    rca_cols = ['number', 'question', 'assessment', 'reason']
    for shape in slide.shapes:
        if not shape.has_table:
            continue
        table = shape.table
        header = table_header_text(table)
        if 'rca question' in header or ('#' in header and 'risk' in header):
            rows = table_to_rows(table, rca_cols)
            if rows:
                fields['rca_rows'] = rows
        elif 'description' in header and 'probability' in header:
            rows = table_to_rows(table, RISK_COLS)
            if rows:
                fields['risk_rows'] = rows


def _extract_staffing_approvals(slide, fields):
    loc_cols = ['location', 'band_mix', 'staffing_pct', 'approver', 'approval_date']
    for shape in slide.shapes:
        if not shape.has_table:
            continue
        table = shape.table
        header = table_header_text(table)
        if 'location' in header and 'band mix' in header:
            rows = table_to_rows(table, loc_cols)
            if rows:
                fields['Delivery Locations'] = rows
        elif 'description' in header and 'probability' in header:
            rows = table_to_rows(table, RISK_COLS)
            if rows:
                fields['Key Resourcing Risks'] = rows


def _extract_ibm_consulting_lines(slide, fields):
    for shape in slide.shapes:
        if not shape.has_table:
            continue
        table = shape.table
        header = table_header_text(table)
        if 'service line' in header and 'sub-line' in header:
            cols = ['col_0', 'service_line', 'sub_line', 'tcv_usd']
            rows = table_to_rows(table, cols)
            if rows:
                fields['IBM Consulting Lines'] = rows
            return


def _extract_executive_summary(slide, fields):
    req_idx = 0
    for shape in slide.shapes:
        if shape.has_table:
            table = shape.table
            header = table_header_text(table)
            if 'customer requirement' in header or ('requirement' in header and 'solution' in header):
                for row in list(table.rows)[1:]:
                    cells = [c.text.strip() for c in row.cells]
                    if not any(cells):
                        continue
                    prefix = f'exec_req_{req_idx}'
                    if len(cells) > 0:
                        fields[f'{prefix}_requirement'] = cells[0]
                    if len(cells) > 1:
                        fields[f'{prefix}_provider'] = cells[1]
                    if len(cells) > 2:
                        fields[f'{prefix}_l30'] = cells[2]
                    if len(cells) > 3:
                        fields[f'{prefix}_scope'] = cells[3]
                    req_idx += 1
        elif shape.has_text_frame:
            text = shape.text_frame.text.strip()
            if not text or len(text) < 30:
                continue
            name_lower = (shape.name or '').lower()
            if 'ask' in name_lower or 'client' in name_lower:
                fields['exec_client_ask'] = text
            elif 'win' in name_lower or 'theme' in name_lower:
                fields['exec_win_themes'] = text
            elif 'solution' in name_lower or 'summary' in name_lower:
                fields['exec_solution_summary'] = text


def extract_isbd_fields(pptx_path):
    """Extract all iSBD fields from a PPTX.
    Returns (fields, slide_titles, is_valid, error_message)."""
    try:
        from pptx import Presentation
    except ImportError:
        return {}, [], False, 'python-pptx not installed'
    try:
        prs = Presentation(pptx_path)
    except Exception as exc:
        return {}, [], False, f'Cannot open PPTX: {exc}'

    fields = {}
    slide_titles = []

    for slide in prs.slides:
        title_text = get_slide_title(slide)
        if title_text:
            slide_titles.append(title_text)

        slide_type = SLIDE_TITLE_MAP.get(title_text.lower().strip()) if title_text else None

        _extract_kv_from_slide(slide, fields, slide_type)

        if slide_type == 'solution_history':
            _extract_solution_history(slide, fields)
        elif slide_type == 'executive_summary':
            _extract_executive_summary(slide, fields)
        elif slide_type == 'risk_assessment':
            _extract_risk_assessment(slide, fields)
        elif slide_type == 'compliance_review':
            _extract_compliance_review(slide, fields)
        elif slide_type == 'staffing_approvals':
            _extract_staffing_approvals(slide, fields)
        elif slide_type == 'deal_profile_summary':
            _extract_ibm_consulting_lines(slide, fields)

    title_set = {t.lower().strip() for t in slide_titles}
    matches = title_set & ISBD_SIGNATURE_TITLES
    is_valid = len(matches) >= ISBD_MIN_TITLE_MATCHES

    if not is_valid:
        msg = f'File does not appear to be an iSBD. Found {len(matches)}/{ISBD_MIN_TITLE_MATCHES} expected section titles.'
        return fields, slide_titles, False, msg

    return fields, slide_titles, True, None
