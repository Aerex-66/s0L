import os

from pptx import Presentation
from pptx.util import Inches

from core.features.documents.isbd_export import export_isbd_pptx
from core.features.documents.isbd_extract import extract_isbd_fields
from core.features.documents.records import apply_isbd_import

PROFILE_ROWS = [
    ('Bid Submission Date', '2026-11-01'),
    ('Opportunity Name', 'Proj X'),
    ('TCV / NPV GP%', '$5M'),
]


def _deck(path):
    prs = Presentation()
    for title, rows in (('Deal Profile Summary', PROFILE_ROWS),
                        ('Executive Summary', [('Win Theme & Negotiation Strategy', 'Be bold')])):
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        slide.shapes.title.text = title
        table = slide.shapes.add_table(len(rows), 2, Inches(1), Inches(1.5), Inches(6), Inches(2)).table
        for r, (label, value) in enumerate(rows):
            table.cell(r, 0).text = label
            table.cell(r, 1).text = value
    prs.save(path)
    return str(path)


def _cell_after(pptx_path, label):
    for slide in Presentation(pptx_path).slides:
        for shape in slide.shapes:
            if shape.has_table:
                for row in shape.table.rows:
                    if row.cells[0].text.strip() == label:
                        return row.cells[1].text
    return None


def test_extract_uses_index_keys(tmp_path):
    fields, _titles, ok, msg = extract_isbd_fields(_deck(tmp_path / 'isbd.pptx'))
    assert ok, msg
    assert fields['bid_submission_date'] == '2026-11-01'
    assert fields['project_name'] == 'Proj X'
    assert fields['tcv'] == '$5M'
    assert fields['win_theme'] == 'Be bold'


def test_import_routes_values_to_their_owners():
    deal = {'id': '1', 'project_name': 'Existing', 'opportunity_owner': '', 'isbd_data': {'swb_stage': '2'}}
    apply_isbd_import(deal, {
        'bid_submission_date': '2026-11-01', 'project_name': 'Deck name', 'opportunity_owner': 'Jane',
        'tcv': '$5M', 'Compliance': [{'compliance_item': 'DPA'}],
    })
    assert deal['isbd_data'] == {
        'swb_stage': '2', 'bid_submission_date': '2026-11-01', 'Compliance': [{'compliance_item': 'DPA'}],
    }
    assert deal['project_name'] == 'Existing'
    assert deal['opportunity_owner'] == 'Jane'


def test_export_writes_value_under_index_key(tmp_path):
    out = export_isbd_pptx(_deck(tmp_path / 'isbd.pptx'), {'bid_submission_date': '2026-12-01'}, timelines=[])
    assert _cell_after(out, 'Bid Submission Date') == '2026-12-01'


def test_export_with_no_financials(tmp_path):
    out = export_isbd_pptx(_deck(tmp_path / 'isbd.pptx'), {}, timelines=[])
    assert os.path.isfile(out)
    assert _cell_after(out, 'TCV / NPV GP%') == '$5M'
